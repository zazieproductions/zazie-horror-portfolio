#!/usr/bin/env python3
"""Build the Zazie Productions / Zazie Kanwar-Torge media master workbook."""

from __future__ import annotations

import csv
from collections import Counter
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "deliverables"
OUT_DIR.mkdir(exist_ok=True)
OUT_FILE = OUT_DIR / "Zazie_Productions_Media_Master.xlsx"
CSV_FILE = OUT_DIR / "Zazie_Productions_Media_Master.csv"
CHECKED = "2026-08-09"

# Palette
INK = "171216"
BURGUNDY = "6F1523"
BLOOD = "98263A"
EMBER = "C75B45"
BONE = "F3ECE5"
PARCHMENT = "FBF8F4"
MIST = "D7CFCA"
ASH = "6D6564"
PALE_RED = "F6E4E7"
PALE_GOLD = "F7EED8"
PALE_GREEN = "E4F0E9"
PALE_BLUE = "E6EEF5"
WHITE = "FFFFFF"
GRAY = "E9E5E2"
DARK_GRAY = "3C3538"

MASTER_HEADERS = [
    "ID", "Priority", "Group", "Media Type", "Published", "Year", "Title / Item",
    "Publication / Platform", "Exact Name Matched", "How Featured", "Work / Subject",
    "Byline / Host", "Language", "Source Relationship", "Verification", "Confidence",
    "Press-Kit Use", "Direct Link", "Secondary Link", "Evidence / Notes", "Last Checked"
]

rows: list[dict] = []


def add(
    id: str,
    priority: str,
    group: str,
    media_type: str,
    published: str,
    title: str,
    publication: str,
    exact: str,
    feature: str,
    subject: str,
    link: str,
    *,
    byline: str = "",
    language: str = "English",
    relationship: str = "Independent editorial",
    verification: str = "Direct page verified",
    confidence: str = "High",
    press_use: str = "Recommended",
    secondary: str = "",
    evidence: str = "",
) -> None:
    year = ""
    if published and published != "Unknown":
        try:
            year = int(published[:4])
        except (ValueError, TypeError):
            year = ""
    rows.append({
        "ID": id,
        "Priority": priority,
        "Group": group,
        "Media Type": media_type,
        "Published": published,
        "Year": year,
        "Title / Item": title,
        "Publication / Platform": publication,
        "Exact Name Matched": exact,
        "How Featured": feature,
        "Work / Subject": subject,
        "Byline / Host": byline,
        "Language": language,
        "Source Relationship": relationship,
        "Verification": verification,
        "Confidence": confidence,
        "Press-Kit Use": press_use,
        "Direct Link": link,
        "Secondary Link": secondary,
        "Evidence / Notes": evidence,
        "Last Checked": CHECKED,
    })


# ---------------------------------------------------------------------------
# PRESS, EDITORIAL, PODCASTS, INSTITUTIONAL COVERAGE
# ---------------------------------------------------------------------------
add("PE-001", "A", "Press & Editorial", "Album feature / review", "2026-06-12",
    "Forget ASMR. Zazie Productions Will Rewire Your Whole Nervous System",
    "Limitless Magazine", "Zazie Productions", "Dedicated feature on the album and artist",
    "Anesthesia for the Signal Age",
    "https://limitless-magazine.com/2026/06/12/forget-asmr-zazie-productions-will-rewire-your-whole-nervous-system",
    byline="AMMPRO", evidence="Headline and article repeatedly name Zazie Productions.")
add("PE-002", "A", "Press & Editorial", "Composer profile / feature", "2026-06-12",
    "How an Underground Experimental Musician Became a Go-To Composer for Psychological Horror",
    "BillboardWire", "Zazie Kanwar-Torge; Zazie Productions",
    "Dedicated profile with direct quote and IMDb/demo-reel links", "Psychological-horror scoring",
    "https://billboardwire.com/how-an-underground-experimental-musician-became-a-go-to-composer-for-psychological-horror/",
    byline="Ethan Walker", evidence="Opening identifies Zazie Kanwar-Torge and Zazie Productions exactly.")
add("PE-003", "A", "Press & Editorial", "Artist profile / feature", "2025-12-12",
    "Zazie Productions: The Underground Polymath Redefining Experimental Music",
    "Grammy Weekly", "Zazie Productions", "Dedicated artist profile", "Career overview",
    "https://grammyweekly.com/zazie-productions-the-underground-polymath-redefining-experimental-music/",
    byline="Angelina Joy", evidence="Exact name appears in headline and throughout article.")
add("PE-004", "A", "Press & Editorial", "Music news / single feature", "2025-02-18",
    "ZAZIE PRODUCTIONS Unleash Eclectic Single ‘Can’t Get My Eyes Off You (123)’",
    "HEAVY Magazine", "Zazie Productions", "Dedicated single feature",
    "Can’t Get My Eyes Off You (123)",
    "https://heavymag.com.au/zazie-productions-unleash-eclectic-single-cant-get-my-eyes-off-you-123/",
    byline="HEAVY", evidence="Exact name appears in headline, copy and tag.")
add("PE-005", "A", "Press & Editorial", "Music review", "2025-02-08",
    "Reseña: Zazie Productions y Unwashed Miscreant rompen esquemas con ‘Can’t Get My Eyes Off You (123)’",
    "Indie AM", "Zazie Productions", "Dedicated Spanish-language review",
    "Can’t Get My Eyes Off You (123)",
    "https://indieam.com.mx/2025/02/08/resena-zazie-productions-y-unwashed-miscreant-rompen-esquemas-con-cant-get-my-eyes-off-you-123/",
    byline="Alejandro Balcázar", language="Spanish", evidence="Exact name appears in headline and article body.")
add("PE-006", "A", "Press & Editorial", "Music feature", "2025-03-17",
    "Zazie Productions y el experimento sonoro que desafía la lógica: ‘CAN’T GET MY EYES OFF YOU (123)’",
    "Shock Web Radio", "Zazie Productions", "Dedicated Spanish-language feature",
    "Can’t Get My Eyes Off You (123)",
    "https://shockwebradio.com/musica/zazie-productions-y-el-experimento-sonoro-que-desafia-la-logica-cant-get-my-eyes-off-you-123/",
    byline="Rita Andino", language="Spanish", evidence="Exact name appears in headline, copy, image metadata and tag.")
add("PE-007", "B", "Press & Editorial", "Music roundup", "2025-03-03",
    "Radar: Shape, Samuel James e outros sons que chegaram pelo Groover",
    "Pop Fantasma", "Zazie Productions", "Artist and track included in a multi-artist roundup",
    "Stunning That You’d Care",
    "https://popfantasma.com.br/radar-shape-samuel-james-e-outros-sons-que-chegaram-pelo-groover/",
    language="Portuguese", evidence="Section heading reads ‘ZAZIE PRODUCTIONS, STUNNING THAT YOU’D CARE.’")
add("PE-008", "A", "Press & Editorial", "Film review", "2025-03-24",
    "Spinning in the Wildness: Phantom Requiem", "Lake Ivan Film Journal",
    "Zazie Kanwar-Torge; Zazie Productions", "Full-length critical review",
    "Phantom Requiem",
    "https://lakeivan.substack.com/p/spinning-in-the-wildness-phantom",
    byline="David Finkelstein", evidence="Review names the short as by Zazie Productions and credits filmmaker Zazie Kanwar-Torge.")
add("PE-009", "B", "Press & Editorial", "Music blog / compilation announcement", "2023-11-11",
    "Dark Ambient: Shadowlands 4 Is Out Now!", "Casey Douglass",
    "Zazie Productions", "Named in embedded compilation tracklist", "Shadowlands 4 / Seroquel Coma",
    "https://www.casey-douglass.com/2023/11/dark-ambient-shadowlands-4-is-out-now.html",
    byline="Casey Douglass", evidence="Embedded tracklist shows ‘Zazie Productions - Seroquel Coma.’")
add("PE-010", "A", "Press & Editorial", "Expert commentary", "2025-10-13",
    "What Creators Really Want From Brands In 2025 (And What They’ll Expect In 2026)",
    "Viral Nation", "Zazie Kanwar-Torge; Zazie Productions",
    "Quoted creator and founder in industry report", "Creator-brand partnerships",
    "https://www.viralnation.com/resources/blog/what-creators-want-from-brands",
    evidence="Exact full name and company name introduce a direct quote.")
add("PE-011", "B", "Press & Editorial", "Expert commentary", "2025-10-03",
    "Art PR: Expanding artistic reach through digital exhibitions", "PR ON THE GO",
    "Zazie Kanwar-Torge; Zazie Productions", "Named expert with extended quote",
    "Digital exhibitions and audience growth",
    "https://pronthego.com/pages/blog/art-PR-expanding-artistic-reach-through-digital-exhibitions",
    evidence="Expert heading contains both exact names.")
add("PE-012", "B", "Press & Editorial", "Expert commentary", "2025-10-27",
    "How film festivals can help a film attract media attention and get distribution deals", "PR ON THE GO",
    "Zazie Kanwar-Torge; Zazie Productions", "Named expert with extended quote",
    "Film-festival publicity",
    "https://pronthego.com/pages/blog/indie-film-PR-how-film-festivals-can-help-a-film-attract-media-attention-and-get-distribution-deals",
    evidence="Expert heading identifies Zazie Kanwar-Torge at Zazie Productions.")
add("PE-013", "B", "Press & Editorial", "Expert commentary", "2025-10-29",
    "Sights and Sounds - Aesthetic Identities of Indie Musicians", "PR ON THE GO",
    "Zazie Kanwar-Torge; Zazie Productions", "Named expert / artist-curator quote",
    "Indie music aesthetics",
    "https://pronthego.com/pages/blog/indie-music-PR-sights-and-sounds-aesthetic-identities-of-indie-musicians",
    evidence="Search-indexed expert heading contains both exact names.")
add("PE-014", "B", "Press & Editorial", "Expert commentary", "2025-04-18",
    "This is how rebellious, alternative fashion is making its way back into society", "PR ON THE GO",
    "Zazie Productions", "Named expert with extended quote", "Grunge and whimsigoth fashion",
    "https://pronthego.com/pages/blog/this-is-how-rebellious-alternative-fashion-is-making-its-way-back-into-society",
    evidence="Contributor section is headed by the exact name Zazie Productions.")
add("PE-015", "B", "Press & Editorial", "Expert commentary", "2025-04-19",
    "When Algorithms Imitate Art: The Rise of Ghibli-Style AI and What It Means for Creators",
    "PR ON THE GO", "Zazie Productions", "Named expert with direct quote", "AI, artistic ethics and imitation",
    "https://pronthego.com/pages/blog/when-algorithms-imitate-art-the-rise-of-ghibli-style-ai-and-what-it-means-for-creators",
    verification="Direct page + search index verified", evidence="Indexed contributor heading uses the exact name Zazie Productions.")
add("PE-016", "C", "Press & Editorial", "Press release", "2024-09-12",
    "Zazie Productions Unveils ‘Phantom Requiem’ (2024) – A Haunting Five-Minute Stop-Motion Avant-Garde Masterpiece",
    "PRFree", "Zazie Productions", "Issued press release", "Phantom Requiem",
    "https://news.prfree.org/@zazieproductions/zazie-productions-unveils-phantom-requiem-2024-a-haunting-five-minute-stop-motion-avant-garde-masterpiece-q85vaihvx2hk",
    relationship="Press-release platform / issued by artist", press_use="Context only",
    evidence="Exact name is in headline and issuer field; label as a press release, not independent review.")
add("PE-017", "C", "Press & Editorial", "Press release", "2025-03-17",
    "Zazie Productions Launches The Vanishing Point Syndicate, a Bold New NetLabel, with Dissonance Index Vol. 1",
    "PRFree", "Zazie Productions", "Issued press release", "The Vanishing Point Syndicate / Dissonance Index Vol. 1",
    "https://news.prfree.org/@noisemusicnewsdaily/zazie-productions-launches-the-vanishing-point-syndicate-a-bold-new-netlabel-with-dissonance-index-vol-1-a-global-underground-compilation-krmhmvbx1u59",
    relationship="Press-release platform", press_use="Context only",
    evidence="Exact name appears in headline and release body.")
add("PE-018", "C", "Press & Editorial", "Hosted article / release", "2025-03-17",
    "Zazie Productions Unleashes The Vanishing Point Syndicate with Dissonance Index Compilation",
    "Telegraph (telegra.ph)", "Zazie Productions", "Dedicated hosted article", "The Vanishing Point Syndicate",
    "https://telegra.ph/Zazie-Productions-Unleashes-The-Vanishing-Point-Syndicate-with-Dissonance-Index-Compilation-03-17",
    byline="Pablo Varga", relationship="User-hosted publishing platform", press_use="Context only",
    evidence="Exact name appears in headline and article body. This is Telegraph by Telegram, not The Daily Telegraph.")
add("PE-019", "B", "Press & Editorial", "Podcast episode", "2025-03-18",
    "The Online Frontier of Sound: Exploring The Vanishing Point Syndicate's First Compilation",
    "Five Steps From The Fringe / Spotify for Creators", "Zazie Productions",
    "Episode centers the artist’s new online label", "The Vanishing Point Syndicate",
    "https://creators.spotify.com/pod/profile/fsftf/episodes/The-Online-Frontier-of-Sound-Exploring-The-Vanishing-Point-Syndicates-First-Compilation-e30bl26",
    byline="Obsidian Veil Media", relationship="Third-party podcast", evidence="Episode description begins with the exact name Zazie Productions.")
add("PE-020", "A", "Press & Editorial", "Institutional artist feature / broadcast", "2021-11-16",
    "BROADCAST: BMC RADIO ARTIST ZAZIE PRODUCTIONS",
    "Black Mountain College Museum + Arts Center", "Zazie Productions",
    "Commissioned radio artist bio, broadcast and interview announcement", "Cheaper Impressions / BMC Radio Art",
    "https://www.blackmountaincollege.org/broadcast-bmc-radio-artist-zazie-productions/",
    relationship="Institutional", evidence="Exact name appears in headline, artist bio and event copy.")
add("PE-021", "A", "Press & Editorial", "Podcast episode", "2024-09-26",
    "Phantom Requiem: The Haunting Vision of Zazie Productions", "Five Steps From The Fringe / Spotify",
    "Zazie Productions", "Dedicated podcast episode", "Phantom Requiem",
    "https://open.spotify.com/episode/1fvq4kF8w6ew8IllOLf5yn",
    byline="Obsidian Veil Media", relationship="Third-party podcast", evidence="Exact name is in episode title and description.")
add("PE-022", "A", "Press & Editorial", "Audio feature", "2021-12-17",
    "BMC Radio Art: Zazie Productions - Cheaper Impressions", "Black Mountain College Radio / SoundCloud",
    "Zazie Productions; Zazie Kanwar-Torge", "Hosted commissioned radio artwork and artist description",
    "Cheaper Impressions",
    "https://soundcloud.com/bmcmuseum/bmc-radio-art-zazie-productions-cheaper-impressions",
    relationship="Institutional", evidence="Search-indexed title contains Zazie Productions and description names Zazie Kanwar-Torge.")
add("PE-023", "B", "Press & Editorial", "Podcast listing", "2021-12-17",
    "BMC Radio Art: Zazie Productions - Cheaper Impressions", "Apple Podcasts",
    "Zazie Productions", "Podcast distribution listing", "Cheaper Impressions",
    "https://podcasts.apple.com/au/podcast/bmc-radio-art-zazie-productions-cheaper-impressions/id1538618362?i=1000545099810",
    relationship="Podcast distributor", verification="Region-restricted page; exact title verified by URL and linked source",
    confidence="Medium", evidence="Apple page is region restricted; exact episode title is preserved in URL and linked by official Linktree.")
add("PE-024", "B", "Press & Editorial", "Magazine feature / sound art", "2024-06-18",
    "Superpresent Vol. 4 No. 3 — Suture Disruption (Percussive)", "Superpresent Magazine",
    "Zazie Productions", "Contributor and featured sound-art entry on p. 66", "Suture Disruption (Percussive)",
    "https://superpresent.org/wp-content/uploads/2024/06/V4N3-pdf.pdf",
    relationship="Independent arts magazine", evidence="PDF contents and entry credit the exact name Zazie Productions.")
add("PE-025", "B", "Press & Editorial", "Magazine feature / music", "2024",
    "Issue #4, 2024 — Frostbite Fantasia", "Ranger Magazine", "Zazie Productions",
    "Named music contributor with dedicated page", "Frostbite Fantasia",
    "https://www.rangermagazine.net/zazie_issue4",
    secondary="https://www.rangermagazine.net/issue4",
    relationship="Independent arts magazine", evidence="Dedicated feature page and issue contents use the exact name.")

# ---------------------------------------------------------------------------
# FILM, FESTIVALS, EXHIBITIONS
# ---------------------------------------------------------------------------
add("FF-001", "A", "Film, Festivals & Exhibitions", "Film review / exhibition page", "2024-12",
    "PHANTOM REQUIEM by Zazie Kanwar-Torge (Zazie Productions) | USA (2024) – JURY SPECIAL MENTION",
    "Pebbles Underground", "Zazie Kanwar-Torge; Zazie Productions",
    "Film page, award citation, bio and full film", "Phantom Requiem",
    "https://pebblesunderground.art/video/phantom-requiem/",
    relationship="Independent festival / arts platform", evidence="Both exact names appear in title and credits.")
add("FF-002", "A", "Film, Festivals & Exhibitions", "Award announcement", "2024",
    "AWARD WINNERS – WINTER SCREENINGS 2024", "Pebbles Underground",
    "Zazie Kanwar-Torge; Zazie Productions", "Jury Special Mention", "Phantom Requiem",
    "https://pebblesunderground.art/award-winners-winter-screenings-2024/",
    relationship="Independent festival / arts platform", evidence="Award list uses both exact names.")
add("FF-003", "B", "Film, Festivals & Exhibitions", "Award screening page", "2025",
    "2024 AWARD WINNING WORKS", "Pebbles Underground",
    "Zazie Kanwar-Torge; Zazie Productions", "Award-winning work and creator credits", "Phantom Requiem",
    "https://pebblesunderground.art/2024awardscreening/",
    relationship="Independent festival / arts platform", evidence="Film listing and credit use both exact names.")
add("FF-004", "B", "Film, Festivals & Exhibitions", "Screening page", "2025-09-20",
    "THE FREEDOM TO CONTROL – Pebbles Underground at Visualcontainer [.BOX]", "Pebbles Underground",
    "Zazie Kanwar-Torge; Zazie Productions", "Screening selection and creator credits", "Phantom Requiem",
    "https://pebblesunderground.art/pebblesundergroundvisualcontainerbox/",
    relationship="Independent festival / arts platform", evidence="Screening page includes exact director and alias credits.")
add("FF-005", "A", "Film, Festivals & Exhibitions", "Screening / event listing", "2025-09-20",
    "The Freedom to Control", "Experimental Cinema / Visualcontainer [.BOX]",
    "Zazie Kanwar-Torge; Zazie Productions", "Selected filmmaker and film", "Phantom Requiem",
    "https://expcinema.org/site/en/events/freedom-control",
    relationship="Independent cinema institution", evidence="Program lists ‘PHANTOM REQUIEM by Zazie Kanwar-Torge (Zazie Productions).’")
add("FF-006", "A", "Film, Festivals & Exhibitions", "Award announcement", "2025-01-15",
    "PEBBLES UNDERGROUND SUMMER 2024 AWARD WINNERS + WINTER 2024 AWARD WINNERS",
    "Visualcontainer TV", "Zazie Kanwar-Torge; Zazie Productions",
    "Jury Special Mention and creator listing", "Phantom Requiem",
    "https://www.visualcontainer.tv/pebbles-underground-summer-2024-award-winners-winter-2024-award-winners/",
    relationship="Video-art platform", evidence="Exact full and company names appear in winners list.")
add("FF-007", "A", "Film, Festivals & Exhibitions", "Television festival program", "2024-10-09",
    "Celebrate Halloween with our Very Spooky Film Festival!", "Latest TV Brighton",
    "Zazie Productions", "Film selected for televised festival program", "Phantom Requiem",
    "https://thelatest.co.uk/brighton/2024/10/09/celebrate-halloween-with-our-very-spooky-film-festival/",
    relationship="Independent TV / media outlet", evidence="Program lists ‘Phantom Requiem by Zazie Productions.’")
add("FF-008", "A", "Film, Festivals & Exhibitions", "Gallery exhibition page", "2024",
    "Sepsis — Ripples", "New Media Artspace", "Zazie Productions",
    "Dedicated exhibition page, video, project statement and artist bio", "Sepsis",
    "https://newmediartspace.info/exhibitions/2024_ripples/sepsis.html",
    relationship="Academic / arts exhibition", evidence="Exact name is the artist credit and appears in the bio.")
add("FF-009", "B", "Film, Festivals & Exhibitions", "Group exhibition artist bio", "Unknown",
    "The Artists — Infinite Self Pavilion", "Cybernetic Futures",
    "Zazie Productions", "Named artist with bio in group exhibition", "Infinite Self Pavilion",
    "https://www.cyberneticfutures.com/infinite-self-pavilion/artistsbios",
    relationship="Curated art exhibition", evidence="Artist section is headed ‘ZAZIE PRODUCTIONS.’")
add("FF-010", "A", "Film, Festivals & Exhibitions", "Screening / filmmaker listing", "2025-11-21",
    "Films for Freedom, Presented by The Film-Makers' Cooperative and Canyon Cinema",
    "The Film-Makers' Cooperative", "Zazie Kanwar-Torge",
    "Featured filmmaker in 100+ artist collaborative film", "Films for Freedom",
    "https://film-makerscoop.com/screenings/films-for-freedom-presented-by-the-film-makers-cooperative-a",
    relationship="Film institution", evidence="Featured filmmaker list contains the exact full name.")
add("FF-011", "B", "Film, Festivals & Exhibitions", "Film / video listing", "2025",
    "FILMS FOR FREEDOM", "Canyon Cinema Foundation / Vimeo", "Zazie Kanwar-Torge",
    "Participating filmmaker in hosted collaborative film", "Films for Freedom",
    "https://vimeo.com/1139013273",
    relationship="Film institution / video platform", evidence="Video description includes exact full name in participating filmmakers.")

# ---------------------------------------------------------------------------
# PUBLICATIONS AND RECOGNITION
# ---------------------------------------------------------------------------
add("PR-001", "A", "Publications & Recognition", "Security acknowledgement", "2025-08",
    "Apple web server security acknowledgements — August 2025", "Apple Support",
    "Zazie Kanwar-Torge", "Credited security reporter", "Apple web-server security",
    "https://support.apple.com/en-us/102774",
    relationship="Corporate security acknowledgement", evidence="Apple’s August 2025 credits list contains the exact full name.")
add("PR-002", "A", "Publications & Recognition", "Published letter / finalist profile", "2020",
    "Zazie Kanwar-Torge, Finalist, Local Letters for Global Change", "Pulitzer Center",
    "Zazie Kanwar-Torge", "Published contest-finalist letter and author bio",
    "Letter on LGBTQIA+ rights to Senator Thom Tillis",
    "https://pulitzercenter.org/zazie-kanwar-torge-finalist-local-letters-global-change",
    relationship="Journalism / education institution", evidence="Exact full name is in page title, letter signature and bio.")
add("PR-003", "A", "Publications & Recognition", "Contest announcement", "2020-01-08",
    "Winners and Finalists: Local Letters for Global Change 2019", "Pulitzer Center",
    "Zazie Kanwar-Torge", "Named middle-school finalist", "Local Letters for Global Change 2019",
    "https://pulitzercenter.org/blog/winners-and-finalists-local-letters-global-change-2019",
    relationship="Journalism / education institution", evidence="Finalist list contains the exact full name.")
add("PR-004", "B", "Publications & Recognition", "Anthology contributor listing", "2026-04-30",
    "Are you really awake?", "Amazon Books", "Zazie Kanwar-Torge",
    "Named contributing story author", "Psychological-horror anthology",
    "https://www.amazon.com/Are-really-awake-Aiden-Messer/dp/B0GSVBXBGS",
    secondary="https://www.amazon.com/Are-really-awake-Aiden-Messer-ebook/dp/B0GHN8FBP9",
    relationship="Book retailer / publication listing", evidence="Print and Kindle descriptions list the exact full name among contributors.")
add("PR-005", "B", "Publications & Recognition", "Anthology catalog listing", "2026",
    "Are you really awake? — author/publisher catalog", "Smashwords", "Zazie Kanwar-Torge",
    "Named contributor in book description", "Psychological-horror anthology",
    "https://www.smashwords.com/profile/view/Aiden_E.Messer890",
    relationship="Book distribution platform", evidence="Catalog description contains the exact full name.")
add("PR-006", "B", "Publications & Recognition", "Research / sighting record", "2025-01-26",
    "Cobalt-indigo flash followed by matte triangular object moving in abrupt jumps with edge illumination",
    "Enigma Labs", "Zazie Productions", "Report accredited to the multimedia artist",
    "Enigma sighting #311344",
    "https://enigmalabs.io/sighting/311344",
    relationship="Research / reporting platform", verification="Search index verified; live page is client-rendered",
    confidence="High", evidence="Indexed record says ‘accredited to multimedia artist known as Zazie Productions.’")

# ---------------------------------------------------------------------------
# PROFILES, DATABASES AND CATALOGS
# ---------------------------------------------------------------------------
profile_data = [
    ("PF-001","A","Film / credit database","Zazie Kanwar-Torge","IMDb","Zazie Kanwar-Torge; Zazie Productions","Composer, producer and sound-department profile","https://www.imdb.com/name/nm17333332/","Film credits and biography","Industry database"),
    ("PF-002","A","Professional profile","Zazie Kanwar-Torge - Independent Film Composer","LinkedIn","Zazie Kanwar-Torge; Zazie Productions","Professional profile, publications, projects and awards","https://www.linkedin.com/in/zazie-kanwar-torge-3b8a98373/","Career and publication index","Professional network"),
    ("PF-003","B","Film / creative-industry profile","Zazie Kanwar-Torge","Stage 32","Zazie Kanwar-Torge; Zazie Productions","Biography, awards and film credits","https://www.stage32.com/profile/1164424/about","Film and music credits","Industry network"),
    ("PF-004","B","Film / festival profile","Zazie Productions","FilmFreeway","Zazie Productions","Filmography, 10 awards/selections and news links","https://filmfreeway.com/ZazieProductions","Filmography and awards","Festival platform"),
    ("PF-005","B","Professional member profile","Zazie Productions","MUSE","Zazie Productions","Member biography, skills and affiliations","https://museonline.org/profile/zazie-kanwar-torge/","Music and media career","Professional association"),
    ("PF-006","B","Music-services profile / interview","Award Winning Musical Prodigy — Zazie Productions","SoundBetter","Zazie Productions","Profile, nine reviews, credits and interview","https://soundbetter.com/profiles/614805-zazie-productions","Audio-production services","Services marketplace"),
    ("PF-007","C","Film-industry profile","Zazie Productions","WFCN","Zazie Productions","Biography, filmography and review","https://www.wfcn.co/profile/zazie.productions","Film career","Industry network"),
    ("PF-008","C","Musician directory","Musicians Directory - Zazie Productions","Musicians.Directory","Zazie Productions","Composition profile and biography","https://musicians.directory/profile-zazie-productions","Music career","Directory"),
    ("PF-009","C","Musician profile","Zazieproductions + Musician in Asheville, NC","BandMix","Zazie Productions","Biography, instruments and film video","https://www.bandmix.com/zazieproductions/","Music services","Services marketplace"),
    ("PF-010","C","Creative-services profile","Zazie Productions","Casting Call Club","Zazie Productions","Composer, editor, audio engineer, writer and producer profile","https://www.castingcall.club/zazieproductions","Creative services","Industry network"),
    ("PF-011","C","Developer / creator profile","Zazie Productions","itch.io","Zazie Productions","Creator bio, games, tools and digital products","https://zazieproductions.itch.io/","Interactive media","Creator platform"),
    ("PF-012","B","Developer profile","Zazie Kanwar-Torge's Profile","Hackaday.io","Zazie Kanwar-Torge; Zazie Productions","Hardware/audio developer bio and two projects","https://hackaday.io/ZazieProductions","Audio tech and hardware","Developer community"),
    ("PF-013","B","Composer reel","Zazie Productions Reel","ReelCrafter","Zazie Kanwar-Torge; Zazie Productions","Bio, credits and selected scoring reel","https://play.reelcrafter.com/z6vU_hn4T76AE0kv4rtd9w","Film-composition reel","Portfolio platform"),
    ("PF-014","C","Music analytics profile","Zazie Kanwar-Torge","Songstats","Zazie Kanwar-Torge","Cross-platform music analytics profile","https://songstats.com/artist/mljz9dxu/zazie-kanwar-torge","Music analytics","Analytics platform"),
    ("PF-015","B","Music analytics / editorial profile","Zazie Productions - Songs, Events and Music Stats","Viberate","Zazie Productions","Artist overview, tracks, releases and analytics","https://www.viberate.com/artist/zazie-productions/","Music analytics","Analytics platform"),
    ("PF-016","C","Playlist curator profile","Zazie Productions playlist curator on Groover","Groover","Zazie Productions","Verified curator profile and artist-services bio","https://groover.co/en/influencer/profile/0.zazie-productions/","Playlist curation","Music platform"),
    ("PF-017","C","Music profile","Zazie Productions - @ZazieProd on Slaps","Slaps","Zazie Productions","Artist biography, tracks and influences","https://slaps.com/ZazieProd","Music profile","Music platform"),
    ("PF-018","C","Visual-art portfolio","Zazie Productions","Behance","Zazie Productions","Visual portfolio and profile","https://www.behance.net/zaziediya","Graphic and visual art","Portfolio platform"),
    ("PF-019","C","Artist equipment index","Artists Starting with Z","Equipboard","Zazie Productions","Indexed artist / composer listing","https://equipboard.com/pros-index/z","Music gear profile index","Music database"),
    ("PF-020","C","Music streaming artist profile","Zazie Productions","Spotify","Zazie Productions","Streaming artist profile","https://open.spotify.com/artist/4UOgvZEOo7xBhFBjJvlMm0","Discography","Streaming platform"),
    ("PF-021","C","Algorithmic radio playlist","Zazie Productions Radio","Spotify","Zazie Productions","Artist radio playlist and track appearances","https://open.spotify.com/playlist/37i9dQZF1E4wLqQSg8MWKD","Artist radio","Streaming platform"),
    ("PF-022","C","Music streaming artist profile","Zazie Productions on Apple Music","Apple Music","Zazie Productions; Zazie Kanwar-Torge","Albums, singles and composer credits","https://music.apple.com/us/artist/zazie-productions/1623719351","Discography","Streaming platform"),
    ("PF-023","C","Music streaming artist profile","Zazie Productions: albums, songs, concerts","Deezer","Zazie Productions; Zazie Kanwar-Torge","Discography and composer credits","https://www.deezer.com/us/artist/170543657","Discography","Streaming platform"),
    ("PF-024","C","Artist storefront / discography","Music — Zazie Productions","Bandcamp","Zazie Productions","Artist bio and self-released discography","https://zazieproductions.bandcamp.com/","Discography","Artist storefront"),
    ("PF-025","B","Music database profile","Zazie Productions Discography","Discogs","Zazie Productions","Artist releases and 58 appearances / 60 compilation rows","https://www.discogs.com/artist/11354435-Zazie-Productions","Discography and compilation credits","Music database"),
    ("PF-026","C","Music credits profile","Zazie Kanwar- Torge - Credits","Muso.AI","Zazie Kanwar-Torge; Zazie Productions","41 composer credits, lyricist credits and collaborators","https://credits.muso.ai/profile/4010f3b7-9a87-4961-8dfd-917de0ba787e","Verified music credits","Music database"),
    ("PF-027","C","Digital-product storefront","Subscribe to Zazie Productions on Gumroad","Gumroad","Zazie Productions","Creator profile and digital products","https://zazieproductions.gumroad.com/","Digital products","Creator storefront"),
    ("PF-028","C","Link hub / portfolio index","Zazie Productions","Linktree","Zazie Productions; Zazie Kanwar-Torge","Official index of press, profiles, films, releases and products","https://linktr.ee/ZazieProductions","Cross-platform portfolio","Owned profile hub"),
    ("PF-029","C","Music-video profile","Zazie Productions Videos","ReverbNation","Zazie Productions","Artist page and hosted videos","https://www.reverbnation.com/zazieproductions/videos","Music videos","Music platform"),
    ("PF-030","C","Song database entry","I Love Everything About You","Rate Your Music","Zazie Productions","Artist credit and compilation appearance","https://rateyourmusic.com/song/zazie-productions/i-love-everything-about-you/","Late Night Love Letters","Music database"),
    ("PF-031","C","Music storefront profile","Zazie Productions","Amazon Music","Zazie Productions","Artist discography profile","https://www.amazon.com/music/player/artists/B0B14FFGFV/zazie-productions","Discography","Streaming / retail platform"),
]
for pid, pri, mtype, title, pub, exact, feature, link, subject, rel in profile_data:
    add(pid, pri, "Profiles & Catalogs", mtype, "Unknown", title, pub, exact, feature, subject, link,
        relationship=rel, verification="Direct or search-indexed profile verified", press_use="Supplemental",
        evidence=f"Public page or index contains the exact name: {exact}.")

# ---------------------------------------------------------------------------
# DISCOGS MUSIC COMPILATION APPEARANCES (60 currently listed rows)
# ---------------------------------------------------------------------------
music = [
    (1,2022,"Plosives: An Introduction To Ghost Nun","Easy Concessions","Ghost Nun","https://www.discogs.com/release/23533040-Various-Plosives-An-Introduction-To-Ghost-Nun"),
    (2,2023,"Warm Music For Cold Weather: For the Muscular Dystrophy Association","Oh! Alas! Christmas!","New Shagg Plus","https://www.discogs.com/release/29280184-Various-Warm-Music-For-Cold-Weather-For-the-Muscular-Dystrophy-Association"),
    (3,2023,"Shadowlands 4","Seroquel Coma","Owlripper Recordings","https://www.discogs.com/release/28865734-Various-Shadowlands-4"),
    (4,2023,"Christmasasaurus Unleashed!","The Reindeers Are Running","Not On Label","https://www.discogs.com/release/33603147-Various-Christmasasaurus-Unleashed"),
    (5,2023,"Harsh Noise Wall Means Fight Back: Smash Gay Oppression","Cisgender Replacement Therapy (CRT)","Delirium Psychosis Productions","https://www.discogs.com/release/27529131-Various-Harsh-Noise-Wall-Means-Fight-Back-Smash-Gay-Oppression"),
    (6,2023,"Esoterrexus - A Slithering Black Compilation","R'lyeh's Xenolith","Slithering Black Records","https://www.discogs.com/release/28653841-Various-Esoterrexus-A-Slithering-Black-Compilation"),
    (7,2024,"ArrhythNia Digital Net Label Compilation 4: Experimental Music And Noise","The Day That You Killed Yourself","arrhythNia Records","https://www.discogs.com/release/30238304-Various-ArrhythNia-Digital-Net-Label-Compilation-4-Experimental-Music-And-Noise"),
    (8,2024,"42 Seconds #7","Plumbus Overture","Owlripper Recordings","https://www.discogs.com/release/32706429-Various-42-Seconds-7"),
    (9,2025,"One World","Moving in Perpetual Slumber","Petroglyph Music","https://www.discogs.com/release/35745076-Various-One-World"),
    (10,2025,"Thirteenth Quarterly Report Of Argali Records Netlabel: Æther Corruption","Weaponized Apathy","Argali Records","https://www.discogs.com/release/33641268-Various-Thirteenth-Quarterly-Report-Of-Argali-Records-Netlabel-%C3%86ther-Corruption"),
    (11,2025,"Fourteenth Quarterly Report Of Argali Records Netlabel: The Great Reconnection","Pivot And Deflect","Argali Records","https://www.discogs.com/release/34207969-Various-Fourteenth-Quarterly-Report-Of-Argali-Records-Netlabel-The-Great-Reconnection"),
    (12,2025,"Sixteenth Quarterly Report: Our Just Desserts","Vibroscriptorium","Argali Records","https://www.discogs.com/release/35943793-Various-Sixteenth-Quarterly-Report-Our-Just-Desserts"),
    (13,2025,"Exquisite Tones For Oozing Ears - Vol. II","Pyrogenesis","Corn On The Cob Music","https://www.discogs.com/release/33883035-Various-Exquisite-Tones-For-Oozing-Ears-Vol-II"),
    (14,2025,"Echoes Of Ancient Wrath","Anti-Sanity Phase","Dodendans","https://www.discogs.com/release/37767543-Various-Echoes-Of-Ancient-Wrath"),
    (15,2025,"1 Year Anniversary","Dome Collapse","The Elements Of Tech & Bass Recordings","https://www.discogs.com/release/35696575-Various-1-Year-Anniversary"),
    (16,2025,"Experiments on the Witch House","The Høuse FΩrgøt Its Name / Hex Ritual","Inner Demons / Dawn Of Darkness","https://www.discogs.com/release/35531737-Various-Experiments-on-the-Witch-House"),
    (17,2025,"LSD666","Pivot And Deflect","LOUDsilence","https://www.discogs.com/release/34195165-Various-LSD666"),
    (18,2025,"Harsh Noise Corpus Vol. 4 (International Harsh Noise Compilation)","Track credit on grouped master release","The Hills Are Dead Records","https://www.discogs.com/master/4006543-Various-Harsh-Noise-Corpus-Vol4-International-Harsh-Noise-Compilation"),
    (19,2025,"Solidarity: A Benefit Compilation for the ACLU","Pivot and Deflect","Voidstar Productions","https://www.discogs.com/release/33736134-Various-Solidarity-A-Benefit-Compilation-for-the-ACLU"),
    (20,2025,"Beyond The Body","Vox Hemorrhagia","Witch-House","https://www.discogs.com/release/34128706-Various-Beyond-The-Body"),
    (21,2025,"Dissonance Index Vol. 1","Pyrogenesis","The Vanishing Point Syndicate","https://www.discogs.com/release/33890583-Various-Dissonance-Index-Vol-1"),
    (22,2025,"DOSTOIVSKI","No God to witness no crime to prove","Plataforma Recs","https://www.discogs.com/release/34149232-Various-DOSTOIVSKI"),
    (23,2025,"Thinning Veil Compilation","Moving in Perpetual Slumber","Ingrown Records","https://www.discogs.com/release/35567218-Various-Thinning-Veil-Compilation"),
    (24,2025,"Icelock Continuum","Frost Directive XIII (Broadcast Leak ID: BLZRD-741)","Camembert Électrique","https://www.discogs.com/release/36016036-Various-Icelock-Continuum"),
    (25,2025,"Late Night Love Letters","I Love Everything About You","Utopia District","https://www.discogs.com/release/37333392-Various-Late-Night-Love-Letters"),
    (26,2025,"One String","Wire Psalm","Ensemble For Sound Poetry And Contemporary Music","https://www.discogs.com/release/35645056-EFSPACM-One-String"),
    (27,2025,"LSD100","Stokes Diapir","LOUDsilence","https://www.discogs.com/release/32869101-Various-LSD100"),
    (28,2025,"Synthetic Dystopia","Post-Carbon Signal Rot - Comunion V.2.3","The Church Of Noisy Goat","https://www.discogs.com/release/34073494-Various-Synthetic-Dystopia"),
    (29,2025,"Kosmische Lᴧëm - Volume #001","Pivot and Deflect","Spirit of Possibility Music","https://www.discogs.com/release/34852574-Various-Kosmische-L%E1%B4%A7%C3%ABm-Volume-001"),
    (30,2025,"Vampires (A Sounds For The Soul Compilation)","Vlad The Impaler","Sounds for the Soul Records","https://www.discogs.com/release/35511685-Various-Vampires-A-Sounds-For-The-Soul-Compilation"),
    (31,2025,"Noisembryowl 5","Weaponized Apathy","Owlripper Recordings","https://www.discogs.com/release/35850694-Various-Noisembryowl-5"),
    (32,2025,"NYOTGRINDER NOISE COMPILATION vol.2","xJazz-Grind Obstruction (US)","NYOTGRINDER","https://www.discogs.com/release/34846487-Various-NYOTGRINDER-NOISE-COMPILATION-vol2"),
    (33,2025,"Toxic T.","Occupational Hazards","Plataforma Recs","https://www.discogs.com/release/33889353-Various-Toxic-T"),
    (34,2025,"1985","Occupational Hazards","Errant Static","https://www.discogs.com/release/34981730-Various-1985"),
    (35,2025,"Who Is The Dreamer? — A Tribute To David Lynch","Small Transitory Life","The Church Of Noisy Goat","https://www.discogs.com/release/33220155-Various-Who-Is-The-Dreamer-A-Tribute-To-David-Lynch"),
    (36,2025,"7.7","Pivot and Deflect","Not On Label (<1 Self-Released)","https://www.discogs.com/release/33707496-Divergent-Artists-77"),
    (37,2025,"Stonewall/Noisewall Volume 1","Stalactite Of Dead Reckoning","TSHN Productions","https://www.discogs.com/release/35109107-TSHN-Productions-StonewallNoisewall-Volume-1"),
    (38,2025,"Aspirin Age Vol. 5 (A Shoegaze Compilation)","Weaponized Apathy","Broken Sound Tapes","https://www.discogs.com/release/35667886-Various-Aspirin-Age-Vol-5-A-Shoegaze-Compilation"),
    (39,2025,"Tavern Synth Vol. II","The Hearth Keeps Time","Relics Of The Eternal City","https://www.discogs.com/release/36381970-Various-Tavern-Synth-Vol-II"),
    (40,2025,"Ju-On (呪怨) The Music Compilation","Izumi's Disappearance","Dawn Of Darkness","https://www.discogs.com/release/34647799-Various-Ju-On-%E5%91%AA%E6%80%A8-The-Music-Compilation"),
    (41,2025,"Make a Change: A Compilation of Skankin' Pickle Covers","Not Too Late","Rocket Punch Records","https://www.discogs.com/release/34708917-Various-Make-a-Change-A-Compilation-of-Skankin-Pickle-Covers"),
    (42,2026,"Twenty Three Seconds Ov Time Volume 15: Thee Awakening","Galvanized Insides Can Feel No Joy","Autonomous Individuals Network","https://www.discogs.com/release/36703675-Various-Twenty-Three-Seconds-Ov-Time-Volume-15--Thee-Awakening-"),
    (43,2026,"Summoning 1","The Høuse FΩrgøt Its Name / Hex Ritual","Dittany of Crete","https://www.discogs.com/release/36430204-Various-Summoning-1"),
    (44,2026,"Summoning 2","The Spirit’s Bone Colored Air","Dittany of Crete","https://www.discogs.com/release/36959898-Various-Summoning-2"),
    (45,2026,"Summoning 3","The Silence Began Crawling","Dittany of Crete","https://www.discogs.com/release/37565460-Various-Summoning-3"),
    (46,2026,"Field Recording Vol. 3","R09 Recording Water Pond Collage","Gelombang Audiozine","https://www.discogs.com/release/37181856-Various-Field-Recording-Vol-3"),
    (47,2026,"Field Recording Vol. 4","Pond At Night Cicadas / Wildlife Binaural","Gelombang Audiozine","https://www.discogs.com/release/37644762-Various-Field-Recording-Vol-4"),
    (48,2026,"Harsh Noise Corpus Vol. 5","Track credit on compilation","The Hills Are Dead Records","https://www.discogs.com/release/36147124-Various-Harsh-Noise-Corpus-Vol5"),
    (49,2026,"It's Only HNW (But I Like It) IX","Obelisk Of Static Collapse","The Hills Are Dead Records","https://www.discogs.com/release/37511712-Various-Its-Only-HNW-But-I-Like-It-IX"),
    (50,2026,"Two Years","Checksum Failed Successfully","LOUDsilence","https://www.discogs.com/release/36241024-Various-Two-Years"),
    (51,2026,"I Want To Believe (X-files Tribute)","Trust No Signal","Dawn Of Darkness","https://www.discogs.com/release/37900812-VA-I-Want-To-Believe-X-files-Tribute"),
    (52,2026,"Frida Kahlo","Frida and the Broken Column","Plataforma Recs","https://www.discogs.com/release/36350866-Various-Frida-Kahlo"),
    (53,2026,"SHARDS // Vol. 1","Vibroscriptorium","Z-Dimension","https://www.discogs.com/release/36836716-Various-SHARDS--Vol-1"),
    (54,2026,"Frequências Cadavéricas - Volume I (Compilação)","Negative Organ Function","Brutalize Recs","https://www.discogs.com/release/36973218-Various-Frequ%C3%AAncias-Cadav%C3%A9ricas-Volume-I-Compila%C3%A7%C3%A3o"),
    (55,2026,"H.P. Lovecraft’s ~Whisperer In Darkness~","Akeley's Phonograph","GATES of HYPNOS","https://www.discogs.com/release/38101452-Various-HP-Lovecrafts-Whisperer-In-Darkness"),
    (56,2026,"Noise Around The World 14","Pyrogenesis","Plataforma Recs","https://www.discogs.com/release/36641662-Various-Noise-Around-The-World-14"),
    (57,2026,"Music Inspired By Mulholland Drive","Club Silencio Acoustic Survey","Dawn Of Darkness","https://www.discogs.com/release/36906384-Various-Music-Inspired-By-Mulholland-Drive"),
    (58,2026,"Ecstatic Feedback II: A Terminal Future Industries Compilation","Egregore Intrasound","Terminal Future Industries","https://www.discogs.com/release/37824543-Various-Ecstatic-Feedback-II-A-Terminal-Future-Industries-Compilation"),
    (59,2026,"Psych Against Cancer Vol 3 Part 2","Weaponized Apathy","Psych Lovers","https://www.discogs.com/release/36690658-Various-Psych-Against-Cancer-Vol-3-Part-2-"),
    (60,None,"Arboreal Telegraph","Isochamber Drift","Camembert Électrique","https://www.discogs.com/release/38101530-Various-Arboreal-Telegraph"),
]

for n, year, title, track, label, link in music:
    add(f"MC-{n:03d}", "C", "Music Compilations", "Compilation appearance",
        str(year) if year else "Unknown", title, label, "Zazie Productions",
        "Credited artist / track appearance", track, link,
        relationship="Third-party release / Discogs catalog record",
        verification="Verified via exact-name Discogs artist index",
        press_use="Discography appendix",
        evidence="Discogs currently lists this row under the exact-name Zazie Productions artist appearances index.")

# ---------------------------------------------------------------------------
# LEADS THAT DO NOT YET PASS THE STRICT DIRECT-PAGE RULE
# ---------------------------------------------------------------------------
leads = [
    ("L-001","The State of Onboarding 2025","Mailmodo","LinkedIn publication list; direct article URL not located","Search Mailmodo archive / request author link"),
    ("L-002","Anchor and Ascend","100Subtexts Magazine","LinkedIn publication list; direct issue URL not located","Search publication archive / request issue PDF"),
    ("L-003","Dissociation Support Group to Hold Meeting, Attendees Plan to Show Up (In Spirit)","The Squeaky Wheel","Direct article exists but does not visibly render the exact author name; LinkedIn attributes it","Confirm author metadata with publisher"),
    ("L-004","Essentially Isometric Invertibility For Co-countably Unique, Canonically Bounded Systems","Academia.edu","LinkedIn publication list; direct paper URL not located","Request Academia.edu URL"),
    ("L-005","Phantom Requiem","Mande/Maude Literary Journal for Bipolar Talent","LinkedIn/creator-profile claim; direct issue URL not located","Confirm exact publication name and issue"),
    ("L-006","THRESHOLDS (A Micro Fiction Anthology)","SOM","LinkedIn publication list; direct anthology URL not located","Confirm publisher acronym and ISBN/URL"),
    ("L-007","Life in Limbo Magazine feature","Life in Limbo Magazine","Creator-profile claim only; direct exact-name page not found","Search issue PDFs"),
    ("L-008","Radioclick Digital feature","Radioclick Digital","Creator-profile claim only; direct exact-name page not found","Confirm domain and article title"),
    ("L-009","Breaking Hits feature","Breaking Hits","Creator-profile claim only; direct exact-name page not found","Search platform archive / app"),
    ("L-010","Typescript Magazine feature","The Typescript","SoundBetter/Slaps claim only; direct issue URL not found","Request issue or scan"),
    ("L-011","BBC feature","BBC","SoundBetter claim only; no exact-name BBC page found","Request program/title/date"),
    ("L-012","Zazie Productions: The Eccentric Savant Who Is Redefining Sound and Vision","Unidentified ‘Telegraph’ link","Title appears on FilmFreeway but target URL is absent","Do not confuse with the verified telegra.ph article"),
    ("L-013","What Is the Resonant Field Composer? Inside the Rumored Gesture-Based Music Interface from Zazie Productions","Unknown outlet","Title appears on FilmFreeway; no target URL found","Locate original publication"),
    ("L-014","Last.fm Zazie Productions stats","Last.fm","Linked from Linktree but direct exact-name page was not search-verifiable","Check live artist URL manually"),
    ("L-015","YouTube @zazieproductions","YouTube","Linked from FilmFreeway and Linktree; search did not return exact-name channel page","Verify channel display name in browser"),
]

# ---------------------------------------------------------------------------
# WORKBOOK HELPERS
# ---------------------------------------------------------------------------

def set_sheet_title(ws, title: str, subtitle: str, end_col: int) -> int:
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.fill = PatternFill("solid", fgColor=INK)
    c.font = Font(name="Aptos Display", size=22, bold=True, color=BONE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 36
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    c = ws.cell(2, 1, subtitle)
    c.fill = PatternFill("solid", fgColor=BURGUNDY)
    c.font = Font(name="Aptos", size=10, italic=True, color=WHITE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 24
    return 4


def style_table_sheet(ws, headers: list[str], data: list[dict], title: str, subtitle: str,
                      widths: dict[str, int] | None = None, table_name: str = "Table1") -> None:
    start = set_sheet_title(ws, title, subtitle, len(headers))
    for col, header in enumerate(headers, 1):
        cell = ws.cell(start, col, header)
        cell.fill = PatternFill("solid", fgColor=DARK_GRAY)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=BLOOD))
    ws.row_dimensions[start].height = 34

    link_headers = {"Direct Link", "Secondary Link", "Link", "Evidence Link"}
    for r_idx, item in enumerate(data, start + 1):
        for c_idx, header in enumerate(headers, 1):
            val = item.get(header, "")
            cell = ws.cell(r_idx, c_idx, val)
            cell.font = Font(name="Aptos", size=9, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if header in link_headers and isinstance(val, str) and val.startswith("http"):
                cell.hyperlink = val
                cell.style = "Hyperlink"
                cell.font = Font(name="Aptos", size=9, color="1F5B8F", underline="single")
        if r_idx % 2 == 0:
            for c_idx in range(1, len(headers) + 1):
                ws.cell(r_idx, c_idx).fill = PatternFill("solid", fgColor=PARCHMENT)

    end = start + len(data)
    if data:
        ref = f"A{start}:{get_column_letter(len(headers))}{end}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False
        )
        ws.add_table(table)
    ws.freeze_panes = f"A{start + 1}"
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{end}"
    ws.row_dimensions.group(start + 1, end, hidden=False)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_layout_view = False
    ws.sheet_view.zoomScale = 75

    default_widths = {
        "ID": 11, "Priority": 9, "Group": 25, "Media Type": 25, "Published": 13, "Year": 9,
        "Title / Item": 45, "Publication / Platform": 28, "Exact Name Matched": 28,
        "How Featured": 35, "Work / Subject": 32, "Byline / Host": 22, "Language": 12,
        "Source Relationship": 30, "Verification": 32, "Confidence": 12,
        "Press-Kit Use": 20, "Direct Link": 48, "Secondary Link": 40,
        "Evidence / Notes": 50, "Last Checked": 14, "Link": 48, "Label / Publisher": 28,
        "Track / Contribution": 38, "Follow-up Action": 42, "Why Not Counted Yet": 48,
    }
    if widths:
        default_widths.update(widths)
    for idx, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(idx)].width = default_widths.get(header, 18)


def write_master(ws, data: list[dict], title: str, subtitle: str, table_name: str) -> None:
    style_table_sheet(ws, MASTER_HEADERS, data, title, subtitle, table_name=table_name)
    # Conditional priority accents.
    start, end = 5, 4 + len(data)
    pcol = MASTER_HEADERS.index("Priority") + 1
    for r in range(start, end + 1):
        c = ws.cell(r, pcol)
        if c.value == "A": c.fill = PatternFill("solid", fgColor=PALE_GREEN)
        elif c.value == "B": c.fill = PatternFill("solid", fgColor=PALE_GOLD)
        elif c.value == "C": c.fill = PatternFill("solid", fgColor=GRAY)
        c.font = Font(name="Aptos", size=9, bold=True, color=INK)
        c.alignment = Alignment(horizontal="center", vertical="top")


# Create workbook
wb = Workbook()
ws_readme = wb.active
ws_readme.title = "Read Me"
ws_dashboard = wb.create_sheet("Dashboard")
ws_master = wb.create_sheet("Master Index")
ws_press = wb.create_sheet("Press & Editorial")
ws_film = wb.create_sheet("Film & Exhibitions")
ws_recog = wb.create_sheet("Publications & Recognition")
ws_profiles = wb.create_sheet("Profiles & Catalogs")
ws_music = wb.create_sheet("Music Compilations")
ws_leads = wb.create_sheet("Leads to Verify")
ws_excl = wb.create_sheet("Exclusions & Rules")
ws_log = wb.create_sheet("Research Log")

# Read Me
ws = ws_readme
set_sheet_title(ws, "ZAZIE MEDIA MASTER", "Exact-name public-web census • prepared 2026-08-09", 8)
ws.sheet_properties.tabColor = BLOOD
readme_sections = [
    (4, "Purpose", "An organized, filterable master spreadsheet of public media that visibly features either exact string ‘Zazie Productions’ or ‘Zazie Kanwar-Torge.’ It separates independent press from institutional listings, profiles, self-published material, and compilation credits."),
    (7, "Strict inclusion rule", "A counted row must have a public page, indexed page, PDF, audio listing, book listing, database entry, or release record that contains at least one exact target name. Pages containing only unrelated people named Zazie, ‘Zazie Films,’ ‘ZazieProductions’ without a space, or spelling variants alone are not counted."),
    (10, "What “all” means", "This is a best-effort, extensive public-web census as of the last-checked date—not a claim that every unindexed, deleted, paywalled, social-only, print-only, or private item on the internet can be discovered. Leads that did not pass direct verification are preserved on a separate sheet and excluded from counted totals."),
    (13, "How to use", "Start with Dashboard. Filter Master Index by Group, Priority, Source Relationship, Verification, year, or exact name. Priority A is strongest for a press kit; B is useful supporting coverage; C is catalog/profile/discography material. The Music Compilations tab contains Discogs’ current 60 listed compilation rows."),
    (16, "Source relationship matters", "Independent editorial and institutional records are separated from distributor profiles, user-created profiles, hosted press releases, and self-published pages. A listing here proves the exact-name appearance—not the truth of every biographical or promotional claim made on the linked page."),
    (19, "Workbook contents", "Master Index (all counted URL-level items); focused category tabs; Music Compilations; Leads to Verify; Exclusions & Rules; and Research Log. All URLs are clickable. A companion CSV contains the full Master Index."),
]
for r, heading, body in readme_sections:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = ws.cell(r, 1, heading)
    c.fill = PatternFill("solid", fgColor=BURGUNDY)
    c.font = Font(name="Aptos Display", size=13, bold=True, color=WHITE)
    ws.merge_cells(start_row=r+1, start_column=1, end_row=r+1, end_column=8)
    c = ws.cell(r+1, 1, body)
    c.fill = PatternFill("solid", fgColor=PARCHMENT)
    c.font = Font(name="Aptos", size=10, color=INK)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r+1].height = 46
for col in range(1, 9):
    ws.column_dimensions[get_column_letter(col)].width = 18
ws.freeze_panes = "A4"
ws.sheet_view.showGridLines = False

# Dashboard
ws = ws_dashboard
set_sheet_title(ws, "DASHBOARD", "Counts, source mix and fast navigation • strict counted items only", 12)
ws.sheet_properties.tabColor = EMBER
all_count = len(rows)
group_counts = Counter(r["Group"] for r in rows)
priority_counts = Counter(r["Priority"] for r in rows)
relationship_rollup = Counter()
for r in rows:
    rel = r["Source Relationship"]
    if any(x in rel for x in ["Institution", "Corporate", "Journalism", "Film institution", "Academic"]):
        relationship_rollup["Institutional"] += 1
    elif "Independent" in rel or "Third-party" in rel or "Video-art" in rel or "Curated" in rel:
        relationship_rollup["Independent / third-party"] += 1
    elif "Press-release" in rel or "User-hosted" in rel:
        relationship_rollup["Press release / user-hosted"] += 1
    else:
        relationship_rollup["Profile / catalog / distributor"] += 1

kpis = [
    ("TOTAL VERIFIED URL-LEVEL ITEMS", all_count, "All rows in Master Index"),
    ("PRESS & EDITORIAL", group_counts["Press & Editorial"], "Articles, reviews, podcasts and features"),
    ("FILM / FESTIVALS / EXHIBITIONS", group_counts["Film, Festivals & Exhibitions"], "Screenings, awards and gallery pages"),
    ("PUBLICATIONS / RECOGNITION", group_counts["Publications & Recognition"], "Authored works and formal acknowledgements"),
    ("PROFILES / CATALOGS", group_counts["Profiles & Catalogs"], "Industry databases, portfolios and streaming profiles"),
    ("MUSIC COMPILATION ROWS", group_counts["Music Compilations"], "Current Discogs-listed compilation rows"),
    ("PRIORITY A", priority_counts["A"], "Strongest external press-kit material"),
    ("LEADS NOT COUNTED", len(leads), "Separate follow-up queue"),
]
for i, (label, value, note) in enumerate(kpis):
    row = 4 + (i // 4) * 4
    col = 1 + (i % 4) * 3
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
    c = ws.cell(row, col, label)
    c.fill = PatternFill("solid", fgColor=BURGUNDY)
    c.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=row+1, start_column=col, end_row=row+1, end_column=col+1)
    c = ws.cell(row+1, col, value)
    c.fill = PatternFill("solid", fgColor=BONE)
    c.font = Font(name="Aptos Display", size=22, bold=True, color=BLOOD)
    c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=row+2, start_column=col, end_row=row+2, end_column=col+1)
    c = ws.cell(row+2, col, note)
    c.fill = PatternFill("solid", fgColor=PARCHMENT)
    c.font = Font(name="Aptos", size=8, italic=True, color=ASH)
    c.alignment = Alignment(horizontal="center", wrap_text=True)

# Dashboard source data and charts
base_row = 14
ws.cell(base_row, 1, "Group").font = Font(bold=True, color=WHITE)
ws.cell(base_row, 2, "Count").font = Font(bold=True, color=WHITE)
for c in (ws.cell(base_row,1), ws.cell(base_row,2)): c.fill = PatternFill("solid", fgColor=DARK_GRAY)
for i, (g, count) in enumerate(group_counts.items(), base_row+1):
    ws.cell(i,1,g); ws.cell(i,2,count)
chart = BarChart()
chart.type = "bar"
chart.style = 10
chart.title = "Verified items by group"
chart.y_axis.title = "Group"
chart.x_axis.title = "Count"
chart.add_data(Reference(ws, min_col=2, min_row=base_row, max_row=base_row+len(group_counts)), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=1, min_row=base_row+1, max_row=base_row+len(group_counts)))
chart.height = 7.5; chart.width = 13
chart.legend = None
ws.add_chart(chart, "D14")

rel_row = 23
ws.cell(rel_row,1,"Source relationship"); ws.cell(rel_row,2,"Count")
for c in (ws.cell(rel_row,1), ws.cell(rel_row,2)):
    c.font=Font(bold=True,color=WHITE); c.fill=PatternFill("solid",fgColor=DARK_GRAY)
for i,(g,count) in enumerate(relationship_rollup.items(),rel_row+1):
    ws.cell(i,1,g); ws.cell(i,2,count)
pie=PieChart(); pie.title="Source mix"; pie.height=7.2; pie.width=10
pie.add_data(Reference(ws,min_col=2,min_row=rel_row,max_row=rel_row+len(relationship_rollup)),titles_from_data=True)
pie.set_categories(Reference(ws,min_col=1,min_row=rel_row+1,max_row=rel_row+len(relationship_rollup)))
ws.add_chart(pie,"D29")

# Navigation
nav_row = 39
ws.merge_cells(start_row=nav_row, start_column=1, end_row=nav_row, end_column=12)
c=ws.cell(nav_row,1,"WORKBOOK NAVIGATION"); c.fill=PatternFill("solid",fgColor=BURGUNDY); c.font=Font(bold=True,color=WHITE)
for idx, name in enumerate(["Master Index","Press & Editorial","Film & Exhibitions","Publications & Recognition","Profiles & Catalogs","Music Compilations","Leads to Verify","Exclusions & Rules","Research Log"], nav_row+1):
    c=ws.cell(idx,1,name); c.hyperlink=f"#'{name}'!A1"; c.style="Hyperlink"; c.font=Font(color="1F5B8F",underline="single",bold=True)
    ws.merge_cells(start_row=idx,start_column=1,end_row=idx,end_column=3)
ws.sheet_view.showGridLines=False
for col in range(1,13): ws.column_dimensions[get_column_letter(col)].width=14

# Data sheets
write_master(ws_master, rows, "MASTER INDEX", f"{len(rows)} strict, exact-name URL-level records • filter every column", "MasterIndexTable")
ws_master.sheet_properties.tabColor = BLOOD
press_rows = [r for r in rows if r["Group"] == "Press & Editorial"]
film_rows = [r for r in rows if r["Group"] == "Film, Festivals & Exhibitions"]
recog_rows = [r for r in rows if r["Group"] == "Publications & Recognition"]
profile_rows = [r for r in rows if r["Group"] == "Profiles & Catalogs"]
write_master(ws_press, press_rows, "PRESS & EDITORIAL", "Independent coverage, expert commentary, podcasts and clearly labeled hosted releases", "PressTable")
ws_press.sheet_properties.tabColor = EMBER
write_master(ws_film, film_rows, "FILM, FESTIVALS & EXHIBITIONS", "Screenings, selections, awards, broadcasts and gallery pages", "FilmTable")
ws_film.sheet_properties.tabColor = "574E8C"
write_master(ws_recog, recog_rows, "PUBLICATIONS & RECOGNITION", "Authored publications, formal acknowledgements and credited records", "RecognitionTable")
ws_recog.sheet_properties.tabColor = "A47B25"
write_master(ws_profiles, profile_rows, "PROFILES & CATALOGS", "Industry databases, professional profiles, streaming catalogs and portfolio platforms", "ProfilesTable")
ws_profiles.sheet_properties.tabColor = "3D6878"

# Dedicated music sheet
music_headers = ["Row", "Year", "Compilation", "Track / Contribution", "Label / Publisher", "Exact Name Matched", "Verification", "Link", "Last Checked"]
music_rows = []
for n, year, title, track, label, link in music:
    music_rows.append({
        "Row": n, "Year": year or "Unknown", "Compilation": title, "Track / Contribution": track,
        "Label / Publisher": label, "Exact Name Matched": "Zazie Productions",
        "Verification": "Discogs exact-name artist appearance index", "Link": link, "Last Checked": CHECKED
    })
style_table_sheet(ws_music, music_headers, music_rows, "MUSIC COMPILATIONS",
                  "60 rows currently displayed by Discogs under the exact-name artist index; Discogs summarizes these as 58 appearances",
                  table_name="MusicCompilationTable")
ws_music.sheet_properties.tabColor = "4E7455"

# Leads sheet
lead_headers = ["ID", "Claimed Item", "Claimed Publication", "Why Not Counted Yet", "Follow-up Action", "Evidence Link", "Last Checked"]
lead_rows=[]
for lid,item,pub,why,action in leads:
    evidence_link = "https://www.linkedin.com/in/zazie-kanwar-torge-3b8a98373/" if lid in {"L-001","L-002","L-003","L-004","L-005","L-006"} else "https://zazieproductions.itch.io/"
    if lid in {"L-010","L-011"}: evidence_link="https://soundbetter.com/profiles/614805-zazie-productions"
    if lid in {"L-012","L-013"}: evidence_link="https://filmfreeway.com/ZazieProductions"
    if lid in {"L-014","L-015"}: evidence_link="https://linktr.ee/ZazieProductions"
    lead_rows.append({"ID":lid,"Claimed Item":item,"Claimed Publication":pub,"Why Not Counted Yet":why,
                      "Follow-up Action":action,"Evidence Link":evidence_link,"Last Checked":CHECKED})
style_table_sheet(ws_leads, lead_headers, lead_rows, "LEADS TO VERIFY",
                  "Preserved for follow-up but excluded from Dashboard and Master Index totals", table_name="LeadsTable")
ws_leads.sheet_properties.tabColor = "B07B2E"

# Exclusions & Rules
ex_headers=["Rule / Example","Decision","Reason","Example Link"]
ex_rows=[
    {"Rule / Example":"Exact-name requirement","Decision":"INCLUDE","Reason":"Rendered or indexed page contains ‘Zazie Productions’ or ‘Zazie Kanwar-Torge’ exactly.","Example Link":"https://pebblesunderground.art/video/phantom-requiem/"},
    {"Rule / Example":"French singer Zazie / Isabelle de Truchis de Varennes","Decision":"EXCLUDE","Reason":"Different person; neither exact target name appears.","Example Link":"https://www.imdb.com/name/nm0953919/"},
    {"Rule / Example":"Zazie Films (Toronto production company)","Decision":"EXCLUDE","Reason":"Different company; not ‘Zazie Productions.’","Example Link":"https://international.apfc.info/en/maison-de-production/zazie-films/"},
    {"Rule / Example":"Zzazz Productions / ZaZu Productions / similarly named entities","Decision":"EXCLUDE","Reason":"Near-match only; exact target name absent.","Example Link":"https://www.instagram.com/zzazzproductions/"},
    {"Rule / Example":"ZazieProductions (no space) only","Decision":"EXCLUDE unless page also has exact target","Reason":"User requested exact names. A handle alone is not sufficient.","Example Link":"https://www.gamedevmarket.net/asset/galactic-requiem-immersive-3d-sci-fi-battle-soundscape-2-CkUE"},
    {"Rule / Example":"Zazie Kanwar- Torge (extra space) only","Decision":"EXCLUDE unless page also has exact target","Reason":"Variant alone is not an exact match. Muso.AI is counted only because its page also renders exact ‘Zazie Kanwar-Torge’ and ‘Zazie Productions’ elsewhere.","Example Link":"https://credits.muso.ai/profile/4010f3b7-9a87-4961-8dfd-917de0ba787e"},
    {"Rule / Example":"Current Superpresent issue page","Decision":"REPLACED","Reason":"Current issue changes over time and no longer contains the target; the stable 2024 issue PDF is used instead.","Example Link":"https://superpresent.org/wp-content/uploads/2024/06/V4N3-pdf.pdf"},
    {"Rule / Example":"Claimed feature without retrievable exact-name page","Decision":"LEAD ONLY","Reason":"Preserved on Leads to Verify, but excluded from totals until direct verification.","Example Link":"https://zazieproductions.itch.io/"},
    {"Rule / Example":"Self-issued press release","Decision":"INCLUDE + LABEL","Reason":"It is a public exact-name media URL, but it is explicitly marked as press-release/self-issued material so it is not mistaken for independent editorial.","Example Link":"https://news.prfree.org/@zazieproductions/zazie-productions-unveils-phantom-requiem-2024-a-haunting-five-minute-stop-motion-avant-garde-masterpiece-q85vaihvx2hk"},
]
style_table_sheet(ws_excl, ex_headers, ex_rows, "EXCLUSIONS & RULES",
                  "Guardrails that prevent unrelated Zazies and near-match names from entering the master", table_name="ExclusionsTable")
ws_excl.sheet_properties.tabColor = "7E7A78"

# Research log
log_headers=["Date","Action","Query / Source","Outcome","Notes"]
log_rows=[
    {"Date":CHECKED,"Action":"Exact-name web search","Query / Source":"\"Zazie Kanwar-Torge\"","Outcome":"Profiles, institutional references, publications and film listings located","Notes":"Unrelated generic Zazie results were rejected."},
    {"Date":CHECKED,"Action":"Exact-name web search","Query / Source":"\"Zazie Productions\"","Outcome":"Press, profiles, music catalogs, exhibitions and broadcasts located","Notes":"Near-match entities were rejected."},
    {"Date":CHECKED,"Action":"Targeted outlet search","Query / Source":"HEAVY, Indie AM, Shock Web Radio, Pop Fantasma, PR ON THE GO, Viral Nation","Outcome":"Direct articles verified","Notes":"WordPress APIs were used to recover stable Pop Fantasma and Shock URLs/dates."},
    {"Date":CHECKED,"Action":"Film/festival search","Query / Source":"Pebbles Underground, Visualcontainer, Experimental Cinema, Latest TV, Film-Makers’ Cooperative","Outcome":"Film pages, screening programs and award pages verified","Notes":"Both exact names were retained when present."},
    {"Date":CHECKED,"Action":"Institutional verification","Query / Source":"Pulitzer Center, Apple Support, Black Mountain College","Outcome":"Finalist letter, Apple security credit and radio-art feature verified","Notes":"High-value institutional sources."},
    {"Date":CHECKED,"Action":"Music appearance census","Query / Source":"Discogs artist 11354435, superFilter=Appearances, pages 1–3","Outcome":"60 listed compilation rows captured","Notes":"Discogs header summarizes 58 appearances while its paginated table displays 60 rows, due to grouped/master-release handling."},
    {"Date":CHECKED,"Action":"Profile/catalog sweep","Query / Source":"IMDb, Stage32, FilmFreeway, MUSE, SoundBetter, WFCN, streaming services and directories","Outcome":"31 exact-name profile/catalog URLs retained","Notes":"Categorized as supplemental rather than independent press."},
    {"Date":CHECKED,"Action":"Claim audit","Query / Source":"LinkedIn publications; itch.io/CastingCall/SoundBetter feature claims","Outcome":f"{len(leads)} unresolved leads isolated","Notes":"Not counted until direct exact-name evidence is found."},
]
style_table_sheet(ws_log, log_headers, log_rows, "RESEARCH LOG",
                  "Methods and major source sweeps used to build this edition", table_name="ResearchLogTable")
ws_log.sheet_properties.tabColor = "4D5C72"

# Global styling and print setup
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = "Zazie exact-name media master • &D"
    ws.oddFooter.right.text = "Page &P of &N"

# Dashboard/Read Me remain first; set active Dashboard.
wb.active = wb.sheetnames.index("Dashboard")
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(OUT_FILE)

# Companion CSV (Master Index only)
with CSV_FILE.open("w", newline="", encoding="utf-8-sig") as f:
    writer = csv.DictWriter(f, fieldnames=MASTER_HEADERS)
    writer.writeheader()
    writer.writerows(rows)

# Re-open sanity checks.
check = load_workbook(OUT_FILE, read_only=False, data_only=False)
assert check.sheetnames == ["Read Me","Dashboard","Master Index","Press & Editorial","Film & Exhibitions",
                            "Publications & Recognition","Profiles & Catalogs","Music Compilations",
                            "Leads to Verify","Exclusions & Rules","Research Log"]
assert check["Master Index"].max_row == len(rows) + 4
assert check["Music Compilations"].max_row == len(music_rows) + 4
assert len(rows) == len({r["ID"] for r in rows})
assert len(music) == 60
print(f"Built {OUT_FILE.relative_to(ROOT)} with {len(rows)} verified records and {len(leads)} leads")
print(f"Built {CSV_FILE.relative_to(ROOT)}")
